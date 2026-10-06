# -*- coding: utf-8 -*-
"""PPT7 课程模板 v3 —— 最终版

在 pandoc 默认 reference.pptx 基础上完成（全部在模板内实现，渲染即成品）：
1) 画布 16:9 (13.333 x 7.5 in)，全部占位符重排
2) 深色画布 #101820；标题 #78C8F5 34pt；正文 30/28pt
3) 主题色映射：tx1 走深色 → 表格（PowerPoint 内置样式）自动成为浅底卡片 + 深色文字
4) 母版：顶部章节导航条（24pt）+ 页码域 + 分隔线（封面/章节页隐藏）
5) 封面：金色强调条 + 双层装饰圆环 + 课程信息（无媒体依赖）
6) 章节页：金色强调条
"""
import os, re, zipfile
from pptx import Presentation
from pptx.util import Emu, Inches
from pptx.oxml.ns import qn
from lxml import etree

SRC = "reference-default.pptx"
TMP = "_ref3_stage1.pptx"
OUT = "reference-ppt7.pptx"

W, H = 12192000, 6858000
SCALE = W / 9144000.0

COL = dict(bg="101820", text="F3F6F9", muted="B8C3CE", dim="8296A5", blue="78C8F5",
           blue_dark="1F5F85", gold="FFD166", green="67D6B4", red="FF9F87",
           line="2C4256", panel="16222C", ink="243039")

prs = Presentation(SRC)
prs.slide_width = Emu(W)
prs.slide_height = Emu(H)


def scale_shape(shp, k):
    try:
        l, t, w, h = shp.left, shp.top, shp.width, shp.height
    except Exception:
        return
    if None in (l, t, w, h):
        return
    shp.left, shp.top, shp.width, shp.height = int(l * k), int(t * k), int(w * k), int(h * k)


master = prs.slide_masters[0]
for m in prs.slide_masters:
    for shp in m.shapes:
        scale_shape(shp, SCALE)
    for lay in m.slide_layouts:
        for shp in lay.shapes:
            scale_shape(shp, SCALE)

lays = {l.name: l for l in master.slide_layouts}


def place(layout, idx, l, t, w, h):
    for ph in layout.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.left, ph.top, ph.width, ph.height = Inches(l), Inches(t), Inches(w), Inches(h)
            return ph
    return None


def set_ph_text_style(layout, idx, sz, color, algn="l"):
    """给版式占位符指定第 1 级文字样式（字号/颜色/对齐）"""
    ph = place(layout, idx, 0, 0, 0, 0) if False else None
    for p in layout.placeholders:
        if p.placeholder_format.idx != idx:
            continue
        txBody = p._element.find(qn("p:txBody"))
        lst = txBody.find(qn("a:lstStyle"))
        if lst is None:
            lst = etree.SubElement(txBody, qn("a:lstStyle"))
            txBody.remove(lst)
            txBody.insert(1, lst)
        for ch in list(lst):
            lst.remove(ch)
        lst.append(etree.fromstring(
            f'<a:lvl1pPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            f'marL="0" indent="0" algn="{algn}"><a:buNone/>'
            f'<a:defRPr sz="{sz}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            '<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/>'
            '<a:cs typeface="+mn-cs"/></a:defRPr></a:lvl1pPr>'))


L, R = 0.62, 12.71
CW = R - L
# v4：压缩标题带、下移正文起点、加高正文区，为 28pt+ 正文争取约 8% 纵向空间
TT, TH, BT, BH = 0.84, 0.70, 1.60, 5.62

lay = lays["Title Slide"]
place(lay, 0, 0.98, 2.22, 7.4, 1.85)
place(lay, 1, 1.00, 4.30, 7.4, 1.30)
set_ph_text_style(lay, 1, 2800, COL["text"])

lay = lays["Title and Content"]
place(lay, 0, L, TT, CW, TH)
place(lay, 1, L, BT, CW, BH)

lay = lays["Two Content"]
place(lay, 0, L, TT, CW, TH)
place(lay, 1, L, BT, 5.84, BH)
place(lay, 2, 6.87, BT, 5.84, BH)

lay = lays["Comparison"]
place(lay, 0, L, TT, CW, TH)
place(lay, 1, L, BT, 5.84, 0.66)
place(lay, 2, L, BT + 0.70, 5.84, BH - 0.70)
place(lay, 3, 6.87, BT, 5.84, 0.66)
place(lay, 4, 6.87, BT + 0.70, 5.84, BH - 0.70)

lay = lays["Section Header"]
place(lay, 0, 1.15, 3.02, 11.0, 1.45)
place(lay, 1, 1.15, 4.50, 11.0, 0.95)
set_ph_text_style(lay, 1, 2800, COL["muted"])

lay = lays["Title Only"]
place(lay, 0, L, TT, CW, TH)

lay = lays["Content with Caption"]
place(lay, 0, L, TT, CW, TH)
place(lay, 1, 6.87, BT, 5.84, BH)
place(lay, 2, L, BT, 5.84, BH)

lay = lays["Picture with Caption"]
place(lay, 0, L, TT, CW, TH)
place(lay, 1, L, BT, CW, 4.55)
place(lay, 2, L, 6.30, CW, 0.80)

# ============ v5：统一所有版式占位符的字号，杜绝 21pt/18pt/10.5pt 等小字 ============
# pandoc 默认 reference 中，双栏(21/18pt)、对照(18pt)、图文(24pt/10.5pt)等版式
# 的占位符 lstStyle 自带小字号。这里把每一级的 defRPr sz 统一改写为
# 标题 34pt、一级正文 30pt、二/三级 28pt（符合"最小字号 28pt"要求）。
SKIP_IDX = (10, 11, 12)  # 日期 / 页脚 / 页码占位符


def normalize_ph_sizes(lay):
    for ph in lay.placeholders:
        idx = ph.placeholder_format.idx
        if idx in SKIP_IDX:
            continue
        ptype = str(ph.placeholder_format.type or "")
        is_title = ptype.startswith(("TITLE", "CENTER_TITLE")) or idx == 0
        txBody = ph._element.find(qn("p:txBody"))
        if txBody is None:
            continue
        lst = txBody.find(qn("a:lstStyle"))
        if lst is None:
            lst = etree.Element(qn("a:lstStyle"))
            txBody.insert(1, lst)
        for i in range(1, 10):
            sz = 3400 if is_title else (3000 if i == 1 else 2800)
            tag = qn("a:lvl%dpPr" % i)
            pp = lst.find(tag)
            if pp is None:
                pp = etree.SubElement(lst, tag)
            rpr = pp.find(qn("a:defRPr"))
            if rpr is None:
                rpr = etree.SubElement(pp, qn("a:defRPr"))
            rpr.set("sz", str(sz))
            if is_title:
                rpr.set("b", "1")


for _lay in master.slide_layouts:
    normalize_ph_sizes(_lay)

# ================= XML 片段工具 =================
NS_DECL = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
           'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def _run(txt, sz, color, b=0, typeface="Microsoft YaHei"):
    return (f'<a:r><a:rPr lang="zh-CN" sz="{sz}" b="{b}" dirty="0">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="{typeface}"/><a:ea typeface="{typeface}"/>'
            f'<a:cs typeface="{typeface}"/></a:rPr><a:t>{txt}</a:t></a:r>')


def make_sp(sid, name, l, t, w, h, runs_xml, algn="l"):
    return (f'<p:sp {NS_DECL}><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/>'
            '<p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{int(l * 914400)}" y="{int(t * 914400)}"/>'
            f'<a:ext cx="{int(w * 914400)}" cy="{int(h * 914400)}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            '<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0">'
            f'<a:noAutofit/></a:bodyPr><a:lstStyle/><a:p><a:pPr algn="{algn}"/>{runs_xml}</a:p></p:txBody></p:sp>')


def make_bar(sid, name, l, t, w, h, color, alpha=None):
    fill = (f'<a:solidFill><a:srgbClr val="{color}">' +
            (f'<a:alpha val="{alpha}"/>' if alpha else '') + '</a:srgbClr></a:solidFill>')
    return (f'<p:sp {NS_DECL}><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/>'
            '<p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{int(l * 914400)}" y="{int(t * 914400)}"/>'
            f'<a:ext cx="{int(w * 914400)}" cy="{int(h * 914400)}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill}'
            '<a:ln><a:noFill/></a:ln></p:spPr>'
            '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>')


def make_ring(sid, name, l, t, w, h, color, alpha, line_w_pt):
    return (f'<p:sp {NS_DECL}><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/>'
            '<p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{int(l * 914400)}" y="{int(t * 914400)}"/>'
            f'<a:ext cx="{int(w * 914400)}" cy="{int(h * 914400)}"/></a:xfrm>'
            '<a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom><a:noFill/>'
            f'<a:ln w="{int(line_w_pt * 12700)}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{color}"><a:alpha val="{alpha}"/></a:srgbClr>'
            '</a:solidFill><a:round/></a:ln></p:spPr>'
            '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>')


# ================= 母版：导航条 / 页码 / 分隔线 =================
NAV_TEXT = "项目背景　·　新知探究　·　项目任务　·　项目总结"
pagenum_run = ('<a:fld id="{7A2F1B90-3C41-4E2A-9E77-2B5D0A1C4E01}" type="slidenum">'
               f'<a:rPr lang="zh-CN" sz="2400" b="0"><a:solidFill><a:srgbClr val="{COL["dim"]}"/>'
               '</a:solidFill><a:latin typeface="Microsoft YaHei"/>'
               '<a:ea typeface="Microsoft YaHei"/></a:rPr><a:t>1</a:t></a:fld>')

sptree = master.element.find(qn("p:cSld")).find(qn("p:spTree"))
for xml in (
    make_sp(910, "PPT7-NAV", 3.2, 0.14, 9.51, 0.5,
            _run(NAV_TEXT, 2400, COL["dim"]), algn="r"),
    make_sp(911, "PPT7-PAGE", 0.62, 0.14, 1.6, 0.5, pagenum_run, algn="l"),
    make_bar(912, "PPT7-LINE", 0.62, 0.74, 12.09, 0.014, COL["line"], 70000),
):
    sptree.append(etree.fromstring(xml))

# ================= 标题页版式 =================
title_lay = lays["Title Slide"]
title_lay.element.set("showMasterSp", "0")
for ph in title_lay.placeholders:
    lst = ph._element.find(qn("p:txBody")).find(qn("a:lstStyle"))
    if lst is not None:
        for pp in lst:
            if pp.get("algn") == "ctr":
                pp.set("algn", "l")

t_tree = title_lay.element.find(qn("p:cSld")).find(qn("p:spTree"))
for xml in (
    make_bar(920, "PPT7-CBAR", 1.00, 1.94, 1.5, 0.085, COL["gold"]),
    make_ring(921, "PPT7-RING1", 9.55, 1.15, 4.6, 4.6, COL["blue"], 20000, 1.6),
    make_ring(922, "PPT7-RING2", 10.75, 4.35, 3.4, 3.4, COL["gold"], 16000, 1.2),
    make_sp(923, "PPT7-INFO", 1.00, 5.82, 9.4, 0.5,
            _run("《R 语言与数据可视化》　大数据专业　2 课时（90 分钟）", 2400, COL["dim"]), algn="l"),
):
    t_tree.append(etree.fromstring(xml))

# ================= 章节页版式 =================
sec_lay = lays["Section Header"]
sec_lay.element.set("showMasterSp", "0")
s_tree = sec_lay.element.find(qn("p:cSld")).find(qn("p:spTree"))
for xml in (
    make_bar(930, "PPT7-SECBAR", 1.18, 2.68, 1.5, 0.085, COL["gold"]),
    make_ring(931, "PPT7-SRING", 10.2, 2.6, 4.3, 4.3, COL["blue"], 14000, 1.5),
):
    s_tree.append(etree.fromstring(xml))
for ph in sec_lay.placeholders:
    lst = ph._element.find(qn("p:txBody")).find(qn("a:lstStyle"))
    if lst is not None:
        for pp in lst:
            if pp.get("algn") == "ctr":
                pp.set("algn", "l")

prs.save(TMP)

# ================= 第二阶段：XML 修补 =================
TITLE_STYLE = ('<p:titleStyle><a:lvl1pPr algn="l" defTabSz="342900" rtl="0" eaLnBrk="1" '
               'latinLnBrk="0" hangingPunct="1"><a:spcBef><a:spcPct val="0"/></a:spcBef><a:buNone/>'
               f'<a:defRPr sz="3400" b="1" kern="1200"><a:solidFill><a:srgbClr val="{COL["blue"]}"/>'
               '</a:solidFill><a:latin typeface="+mj-lt"/><a:ea typeface="+mj-ea"/>'
               '<a:cs typeface="+mj-cs"/></a:defRPr></a:lvl1pPr></p:titleStyle>')


def body_lvl(i, sz, marl, bullet, color):
    bu = f'<a:buFont typeface="Arial"/><a:buChar char="{bullet}"/>' if bullet else '<a:buNone/>'
    return (f'<a:lvl{i}pPr marL="{marl}" indent="-342900" algn="l" defTabSz="342900" rtl="0" '
            f'eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
            f'<a:lnSpc><a:spcPct val="95000"/></a:lnSpc>'
            f'<a:spcBef><a:spcPct val="12000"/></a:spcBef>'
            f'{bu}<a:defRPr sz="{sz}" kern="1200"><a:solidFill><a:srgbClr val="{color}"/>'
            f'</a:solidFill><a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/>'
            f'<a:cs typeface="+mn-cs"/></a:defRPr></a:lvl{i}pPr>')


BODY_STYLE = ('<p:bodyStyle>' + body_lvl(1, 3000, 342900, "•", COL["text"])
              + body_lvl(2, 2800, 685800, "‒", COL["text"])
              + "".join(body_lvl(i, 2800, 342900 * (i - 1), "•", COL["text"])
                        for i in (3, 4, 5, 6, 7, 8, 9))
              + '</p:bodyStyle>')

OTHER_STYLE = ('<p:otherStyle><a:defPPr><a:defRPr lang="en-US"/></a:defPPr>'
               + "".join(
                   f'<a:lvl{i}pPr marL="{342900 * (i - 1)}" algn="l" defTabSz="342900" rtl="0" '
                   f'eaLnBrk="1" latinLnBrk="0" hangingPunct="1"><a:defRPr sz="2800" kern="1200">'
                   f'<a:solidFill><a:srgbClr val="{COL["ink"]}"/></a:solidFill>'
                   f'<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/>'
                   f'<a:cs typeface="+mn-cs"/></a:defRPr></a:lvl{i}pPr>' for i in range(1, 10))
               + '</p:otherStyle>')


def patch_theme(xml):
    # 注意：dk1(tx1) 取深色，使内置表格样式成为"浅底 + 深字"；
    # 幻灯片标题/正文颜色由 txStyles 显式指定为浅色，不受影响。
    scheme = ('<a:clrScheme name="CourseDark">'
              f'<a:dk1><a:srgbClr val="{COL["ink"]}"/></a:dk1>'
              f'<a:lt1><a:srgbClr val="{COL["bg"]}"/></a:lt1>'
              f'<a:dk2><a:srgbClr val="{COL["blue_dark"]}"/></a:dk2>'
              f'<a:lt2><a:srgbClr val="{COL["panel"]}"/></a:lt2>'
              f'<a:accent1><a:srgbClr val="{COL["blue"]}"/></a:accent1>'
              f'<a:accent2><a:srgbClr val="{COL["gold"]}"/></a:accent2>'
              f'<a:accent3><a:srgbClr val="{COL["green"]}"/></a:accent3>'
              f'<a:accent4><a:srgbClr val="{COL["red"]}"/></a:accent4>'
              f'<a:accent5><a:srgbClr val="{COL["muted"]}"/></a:accent5>'
              f'<a:accent6><a:srgbClr val="{COL["blue_dark"]}"/></a:accent6>'
              f'<a:hlink><a:srgbClr val="{COL["blue"]}"/></a:hlink>'
              f'<a:folHlink><a:srgbClr val="{COL["muted"]}"/></a:folHlink></a:clrScheme>')
    xml = re.sub(r"<a:clrScheme.*?</a:clrScheme>", scheme, xml, flags=re.S)
    xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*(")',
                 r"\g<1>Microsoft YaHei\g<2>", xml)
    xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*(")',
                 r"\g<1>Microsoft YaHei\g<2>", xml)
    xml = re.sub(r'<a:ea typeface="[^"]*"/>', '<a:ea typeface="Microsoft YaHei"/>', xml)
    xml = re.sub(r'<a:cs typeface="[^"]*"/>', '<a:cs typeface="Microsoft YaHei"/>', xml)
    return xml


def patch_master(xml):
    xml = re.sub(r"<p:bg>.*?</p:bg>", "", xml, flags=re.S)
    bg = (f'<p:bg><p:bgPr><a:solidFill><a:srgbClr val="{COL["bg"]}"/></a:solidFill>'
          f'<a:effectLst/></p:bgPr></p:bg>')
    xml = re.sub(r"(<p:cSld[^>]*>)", lambda m: m.group(1) + bg, xml, count=1)
    xml = re.sub(r"<p:titleStyle>.*?</p:titleStyle>", TITLE_STYLE, xml, flags=re.S)
    xml = re.sub(r"<p:bodyStyle>.*?</p:bodyStyle>", BODY_STYLE, xml, flags=re.S)
    xml = re.sub(r"<p:otherStyle>.*?</p:otherStyle>", OTHER_STYLE, xml, flags=re.S)
    return xml


def patch_layout(xml, name):
    xml = re.sub(r"<p:bg>.*?</p:bg>", "", xml, flags=re.S)
    xml = xml.replace('<a:schemeClr val="tx1"><a:tint val="75000"/></a:schemeClr>',
                      f'<a:srgbClr val="{COL["muted"]}"/>')
    return xml


zin = zipfile.ZipFile(TMP)
zout = zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    n = item.filename
    txt = None
    if n == "ppt/theme/theme1.xml":
        txt = patch_theme(data.decode("utf-8"))
    elif n == "ppt/slideMasters/slideMaster1.xml":
        txt = patch_master(data.decode("utf-8"))
    elif re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", n):
        txt = patch_layout(data.decode("utf-8"), os.path.basename(n))
    zout.writestr(item, txt.encode("utf-8") if txt is not None else data)
zout.close()
zin.close()
os.remove(TMP)
print("reference built:", OUT, os.path.getsize(OUT))
